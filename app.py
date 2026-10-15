import os
import shutil
import json
import time
import base64
import requests
from pathlib import Path
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook

load_dotenv()

API_KEY = os.getenv("GEMINI_API_KEY")

if not API_KEY:
    print("[ERRO] Chave da API nao encontrada. Verifique seu arquivo .env")
    exit()

# -------------------------------------------------------------
# 1. DIRETÓRIOS E PLANILHA DE DESTINO
# -------------------------------------------------------------
PASTA_ENTRADA = Path("entrada")
PASTA_PROCESSADOS = Path("processados")
PASTA_REVISAR = Path("revisar")
ARQUIVO_EXCEL = Path("Controle_Financeiro.xlsx")

PASTA_ENTRADA.mkdir(exist_ok=True)
PASTA_PROCESSADOS.mkdir(exist_ok=True)
PASTA_REVISAR.mkdir(exist_ok=True)


# -------------------------------------------------------------
# 2. EXTRAÇÃO REST COM SUPORTE A PDF E IMAGENS
# -------------------------------------------------------------
def extrair_dados_direto(caminho_arquivo: Path) -> dict:
    extensao = caminho_arquivo.suffix.lower().replace(".", "")
    
    # Mapeamento do MIME Type correto para suportar PDF e imagens
    if extensao == "pdf":
        mime_type = "application/pdf"
    elif extensao in ["jpg", "jpeg"]:
        mime_type = "image/jpeg"
    elif extensao == "png":
        mime_type = "image/png"
    else:
        raise ValueError(f"Formato nao suportado: .{extensao}")

    with open(caminho_arquivo, "rb") as f:
        arquivo_b64 = base64.b64encode(f.read()).decode("utf-8")

    url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-3.8-flash:generateContent?key={API_KEY}"

    payload = {
        "contents": [{
            "parts": [
                {
                    "text": (
                        "Voce e um extrator de dados fiscais e contabeis. "
                        "Analise o documento fornecido (seja PDF ou imagem) e retorne APENAS um JSON valido, "
                        "sem blocos markdown e sem comentarios adicionais, seguindo este formato exato:\n"
                        "{\n"
                        '  "razao_social_emitente": "Nome ou Razao Social",\n'
                        '  "cnpj_cpf_emitente": "00.000.000/0001-00",\n'
                        '  "numero_documento": "12345",\n'
                        '  "data_emissao": "DD/MM/AAAA",\n'
                        '  "valor_total": 0.0,\n'
                        '  "itens": [\n'
                        '    {"descricao": "Nome do item", "quantidade": 1.0, "valor_unitario": 0.0, "valor_total": 0.0}\n'
                        '  ]\n'
                        "}"
                    )
                },
                {"inline_data": {"mime_type": mime_type, "data": arquivo_b64}}
            ]
        }],
        "generationConfig": {
            "temperature": 0.1
        }
    }

    headers = {"Content-Type": "application/json"}

    tentativas = 5
    for i in range(tentativas):
        try:
            resposta = requests.post(url, json=payload, headers=headers)
            dados_resposta = resposta.json()

            if resposta.status_code == 200:
                texto = dados_resposta["candidates"][0]["content"]["parts"][0]["text"].strip()
                
                # Remove delimitadores de codigo se vierem na resposta
                if texto.startswith("```"):
                    texto = texto.split("\n", 1)[-1]
                    if texto.endswith("```"):
                        texto = texto.rsplit("\n", 1)[0]
                        
                return json.loads(texto)
            
            elif resposta.status_code in [503, 429]:
                espera = (i + 1) * 3
                print(f"  [!] Servidor ocupado (Erro {resposta.status_code}). Tentando em {espera}s...")
                time.sleep(espera)
            else:
                raise RuntimeError(f"Erro retornado pela API: {dados_resposta}")

        except Exception as e:
            if i == tentativas - 1:
                raise e
            time.sleep(2)

    raise RuntimeError("Limite de tentativas esgotado sem sucesso na requisicao.")


# -------------------------------------------------------------
# 3. SALVAMENTO NATIVO NO EXCEL (.XLSX) EM MODO APPEND
# -------------------------------------------------------------
def salvar_no_excel(dados: dict, nome_origem: str):
    if not ARQUIVO_EXCEL.exists():
        wb = Workbook()
        ws = wb.active
        ws.title = "Entradas"
        ws.append([
            "Arquivo Origem", "Emitente", "CNPJ/CPF", "Nº Documento", 
            "Data", "Valor Total", "Item", "Quantidade", "Valor Unitario", "Valor Item"
        ])
    else:
        wb = load_workbook(ARQUIVO_EXCEL)
        ws = wb.active

    itens = dados.get("itens", [])
    valor_total_doc = float(dados.get("valor_total") or 0.0)

    if itens:
        for item in itens:
            qtd = float(item.get("quantidade") or 1.0)
            v_unit = float(item.get("valor_unitario") or 0.0)
            v_tot = float(item.get("valor_total") or (qtd * v_unit))

            ws.append([
                nome_origem,
                dados.get("razao_social_emitente", "Nao identificado"),
                dados.get("cnpj_cpf_emitente", "Nao identificado"),
                str(dados.get("numero_documento", "S/N")),
                dados.get("data_emissao", "Nao identificada"),
                valor_total_doc,
                item.get("descricao", "Sem descricao"),
                qtd,
                v_unit,
                v_tot
            ])
    else:
        ws.append([
            nome_origem,
            dados.get("razao_social_emitente", "Nao identificado"),
            dados.get("cnpj_cpf_emitente", "Nao identificado"),
            str(dados.get("numero_documento", "S/N")),
            dados.get("data_emissao", "Nao identificada"),
            valor_total_doc,
            "Sem itens discriminados",
            1.0,
            valor_total_doc,
            valor_total_doc
        ])

    # Autoajuste dinamico da largura de cada coluna para nao cortar textos
    for col in ws.columns:
        comprimento_max = 0
        col_letter = col[0].column_letter
        for cell in col:
            valor = str(cell.value or "")
            if len(valor) > comprimento_max:
                comprimento_max = len(valor)
        ws.column_dimensions[col_letter].width = max(comprimento_max + 3, 12)

    wb.save(ARQUIVO_EXCEL)


# -------------------------------------------------------------
# 4. EXECUÇÃO DO FLUXO PRINCIPAL
# -------------------------------------------------------------
def executar():
    extensoes_permitidas = {".pdf", ".png", ".jpg", ".jpeg"}
    arquivos = [f for f in PASTA_ENTRADA.iterdir() if f.suffix.lower() in extensoes_permitidas]

    if not arquivos:
        print("[!] Nenhum documento localizado na pasta 'entrada/'.")
        print("    Adicione arquivos .pdf, .png ou .jpg para processar.")
        return

    print(f"[*] Iniciando: {len(arquivos)} documento(s) na fila.")

    for arquivo in arquivos:
        print(f"-> Lendo: {arquivo.name}...")
        try:
            dados_extraidos = extrair_dados_direto(arquivo)
            salvar_no_excel(dados_extraidos, arquivo.name)
            
            # Move para processados se o fluxo correu sem erros
            shutil.move(str(arquivo), str(PASTA_PROCESSADOS / arquivo.name))
            print(f"[OK] {arquivo.name} inserido no Excel com sucesso!")
        except Exception as e:
            # Em caso de documento danificado ou erro critico, isola o arquivo
            print(f"[ERRO] Falha ao processar {arquivo.name}: {e}")
            shutil.move(str(arquivo), str(PASTA_REVISAR / arquivo.name))
            print(f"[AVISO] Arquivo movido para '{PASTA_REVISAR.name}/' para conferencia manual.")

        time.sleep(2)

    print(f"\n[CONCLUIDO] Processamento finalizado.")
    print(f"            Planilha atualizada: '{ARQUIVO_EXCEL.resolve()}'.")


if __name__ == "__main__":
    executar()